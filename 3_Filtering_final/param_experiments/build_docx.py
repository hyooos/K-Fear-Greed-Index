# -*- coding: utf-8 -*-
"""필터링 파라미터 검증 리포트를 Word(.docx) 로 생성 — docx/ 폴더에 KO/EN."""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import filter_param_lib as L
import make_notebooks as MN

BASE = Path(__file__).resolve().parent
OUT = BASE / "docx"
OUT.mkdir(exist_ok=True)

VERDICT_ROWS = {
 "ko": [
  ("τ = 0.95  (독성 hard-drop)", "우리 선택", "✓ 검증 — 지표·전략이 τ에 구조적으로 불변. tox≥0.95 댓글의 감성 가중치 비중 0.004%"),
  ("γ = 2  (독성 가중치)", "우리 선택", "✓ 검증 — γ 1·2·3 에서 전략 Sharpe 0.61–0.62, MDD −22.9% 고정. 감쇠 강도의 원칙적 중간값"),
  ("score ≥ 10  (경제 필터 임계값)", "우리 선택", "✓ 검증 — 통과율 cutoff 3–10 = 29.0% 평탄, 13에서 20.0% 급락"),
  ("DF ≥ 1%  (핵심 키워드 기준)", "우리 선택", "◐ 근거 제시 — 0.5%는 일반어 유입, 1%부터 금융어 중심, 1.5%+는 정당 복합어 탈락"),
  ("핵심 10 : 보조 3  (점수 가중치)", "우리 선택", "◐ 근거 제시 — 보조 점수(2·5·10)는 통과 여부에 무관, 순위용"),
  ("핵심 ≥ 1개  (경제 필터 조건)", "우리 선택", "✓ 검증 — 보조 키워드만으로 통과하는 실측 1,165건(통과분 1.2%) 차단"),
  ("2 ~ 6 글자  (키워드 추출)", "방법론 규칙", "✓ 규칙 근거 — 1글자 금융어는 '장'뿐, min 2가 2글자 핵심어 보존 / max 6 vs 8 동일"),
  ("감성 10일 이동평균  (K-FGI 3.4)", "우리 선택", "✓ 검증 — window 5/10/20 중 10이 Sharpe·MDD 모두 최적 (근거가 필요한 값)"),
  ("252일 · 60일 · 70% 분위수", "금융 관행값", "튜닝 대상 아님 — 1년 거래일 / 분기 국면 / 상위 30% 고변동"),
 ],
 "en": [
  ("τ = 0.95  (toxicity hard-drop)", "our choice", "✓ validated — index & strategy structurally invariant to τ; tox≥0.95 weight mass 0.004%"),
  ("γ = 2  (toxicity weighting)", "our choice", "✓ validated — strategy Sharpe 0.61–0.62, MDD −22.9% fixed for γ 1/2/3; principled midpoint"),
  ("score ≥ 10  (economic-filter cutoff)", "our choice", "✓ validated — pass rate flat at 29.0% for cutoffs 3–10, sharp drop to 20.0% at 13"),
  ("DF ≥ 1%  (core-keyword threshold)", "our choice", "◐ justified — 0.5% admits generic words; 1%+ finance-centric; 1.5%+ drops valid compounds"),
  ("core 10 : aux 3  (score weighting)", "our choice", "◐ justified — aux weight (2·5·10) irrelevant to the gate; ranking only"),
  ("core ≥ 1  (economic-filter condition)", "our choice", "✓ validated — removes the 1,165 comments (1.2% of passers) that pass on aux keywords alone"),
  ("2–6 characters  (keyword extraction)", "method rule", "✓ rule-based — only finance 1-char token is '장'; min 2 keeps 2-char terms; max 6 = max 8"),
  ("10-day sentiment MA  (K-FGI §3.4)", "our choice", "✓ validated — window 5/10/20: 10 is optimal on Sharpe and MDD (a value that needs a rationale)"),
  ("252-day · 60-day · 70th pct.", "finance convention", "not tuned — one trading year / quarterly regime / top-30% high volatility"),
 ],
}

DOC_META = {
 "ko": dict(fname="필터링_파라미터_검증_KO.docx",
            title="필터링 파라미터 검증 — 왜 이 값일 수밖에 없었는가",
            subtitle="K-FGI · 3장 필터링 파이프라인",
            verdict_h="검증 결과 요약", verdict_cols=["값", "성격", "판정 · 핵심 근거"],
            concl_from="concl", figword="그림", capword="설명",
            font="Malgun Gothic"),
 "en": dict(fname="Filtering_Parameter_Validation_EN.docx",
            title="Filtering Parameter Validation — why these values, not others",
            subtitle="K-FGI · Ch. 3 Filtering pipeline",
            verdict_h="Summary of verdicts", verdict_cols=["Value", "Type", "Verdict · key evidence"],
            concl_from="concl", figword="Figure", capword="Note",
            font="Calibri"),
}


def set_style_font(doc, name):
    for sn in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        try:
            st = doc.styles[sn]
        except KeyError:
            continue
        st.font.name = name
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        rf.set(qn("w:eastAsia"), name)
    doc.styles["Normal"].font.size = Pt(10.5)


def add_md(doc, text):
    """작은 markdown -> 문단/불릿/굵게/인용."""
    text = re.sub(r"\$w_i = .*?\}\$", "w = (1 − tox)^γ", text)
    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line or line.startswith("|"):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1); continue
        style = None
        if line.startswith("> "):
            line = line[2:]; style = "Intense Quote"
        bullet = line.startswith("- ")
        if bullet:
            line = line[2:]
        p = doc.add_paragraph(style=("List Bullet" if bullet else style))
        # **bold** 분해
        for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", line)):
            r = p.add_run(seg)
            if i % 2 == 1:
                r.bold = True


def add_caption(doc, key, lang):
    Lk = L.LB[lang][key]
    h = doc.add_heading(Lk["t"], level=2)
    img = BASE / "figures" / lang / f"{key}.png"
    doc.add_picture(str(img), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    r = p.add_run(f"[{DOC_META[lang]['capword']}] " + Lk["cap"])
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x44, 0x4B, 0x55)
    p.paragraph_format.space_after = Pt(16)


def build(lang):
    M = DOC_META[lang]; tx = MN.TXT[lang]
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.9)
    set_style_font(doc, M["font"])

    doc.add_paragraph(M["subtitle"]).runs[0].font.color.rgb = RGBColor(0x3A, 0x4D, 0xB8)
    doc.add_heading(M["title"], level=0)

    # 서문
    intro = tx["title"].split("\n", 1)[1].strip()
    for para in intro.split("\n\n"):
        add_md(doc, para)

    # 검증 표
    doc.add_heading(M["verdict_h"], level=1)
    rows = VERDICT_ROWS[lang]
    tbl = doc.add_table(rows=1 + len(rows), cols=3)
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(2.0), Inches(1.05), Inches(3.75)]
    for j, c in enumerate(M["verdict_cols"]):
        cell = tbl.rows[0].cells[j]; cell.text = c
        cell.paragraphs[0].runs[0].bold = True
    for i, (a, b, c) in enumerate(rows, start=1):
        for j, v in enumerate((a, b, c)):
            tbl.rows[i].cells[j].text = v
    for row in tbl.rows:
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    # 섹션 + 그림
    aux_only, aux_pct = 1165, 1.2
    try:
        import pickle
        RR = pickle.loads(L.CACHE.read_bytes())
        aux_only = RR["G"]["aux_only"]; aux_pct = aux_only / RR["G"]["n_pass"] * 100
    except Exception:
        pass

    for sec, keys in MN.SECTIONS:
        body = tx[sec].format(aux_only=aux_only, aux_pct=aux_pct) if "{aux" in tx[sec] else tx[sec]
        add_md(doc, body)
        for k in keys:
            add_caption(doc, k, lang)

    add_md(doc, tx["concl"])
    closing = {
     "ko": "요약하면, τ = 0.95 와 γ = 2 는 감성지표·전략 성과가 그 값에 구조적으로 둔감하여 성능으로 "
           "선택할 수 없는 값이고(따라서 데이터 보존·감쇠 원칙으로 정함), score ≥ 10 은 통과율 평탄 구간의 끝, "
           "감성 10일 이동평균은 위험조정성과의 국소 최적점이다. 상세 판정은 위 표를 참조.",
     "en": "In short, τ = 0.95 and γ = 2 are values the sentiment index and strategy performance are structurally "
           "insensitive to — they cannot have been chosen for performance (they follow from data-preservation and "
           "decay principles); score ≥ 10 is the end of the pass-rate plateau; and the 10-day sentiment MA is the "
           "local optimum on risk-adjusted performance. See the table above for the per-value verdicts.",
    }[lang]
    doc.add_paragraph(closing)

    out = OUT / M["fname"]
    doc.save(str(out))
    print("wrote", out.relative_to(BASE), f"({out.stat().st_size // 1024} KB)")


for lg in ("ko", "en"):
    build(lg)
